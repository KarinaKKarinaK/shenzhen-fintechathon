"""Builds the Zishi pitch deck on the official FinTechathon template.

Reads docs/template/presentation-template.pptx (from fintechathon.g-ican.com/templates/presentation-template.pptx),
keeps its cover, contents and THANKS slides and its layouts, drops its sample slide, and adds the content slides
on the template's content layout (dark banner with the FinTechathon logo, white body).
Outputs docs/Zishi_Pitch_Deck.pptx (speaker notes included) and docs/pdf/Zishi_Pitch_Deck.pdf, drawn from that same
pptx (template backgrounds, every shape at its position) and printed by the system Chrome.
Every number on the slides comes from README.md "By the numbers", docs/EXECUTION_EVIDENCE.md,
SELF_ASSESSMENT.md and web/evidence/fisco/settled-run-*.json.

Run: python web/scripts/deck.py [--png DIR]   (needs python-pptx, pillow, Google Chrome; --png also needs pdftoppm)
"""
import json, subprocess, sys, tempfile
from pathlib import Path
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
TEMPLATE = DOCS / "template/presentation-template.pptx"
W, H = 13.333, 7.5
# template palette: the banner's red, the contents slide's blue 6779C6, white body
RED, INK, GOLD, GREY, CARD, LINE = "c8102e", "1a1a1a", "6779c6", "5c5c5c", "ffffff", "dfe2ee"
FONT = "Arial"
M = 0.75  # side margin, inches
DY = -0.5  # content moved up: the template banner ends at 1.17 in

# cover and contents text
TEAM = "Zishi"
TRACK = "International AI Track, Topic F: Agent Wallet"
DATE = "20 October 2026"
URL = "zishi-fintechathon.vercel.app"
CONTENTS = ["The pilot", "How it works", "Live demo", "Safety", "On chain", "Proof and next"]

# fisco settled run: read the numbers straight from the evidence file
_settled = sorted((ROOT / "web/evidence/fisco").glob("settled-run-*.json"))[-1]
_s = json.loads(_settled.read_text())
_checks = {c["check"]: c for c in _s["checks"]}
WEI = 1e9  # 1 yuan = 1e9 wei in that run
FISCO_CALLS = _checks["every contract call succeeded"]["value"]
LENDER_YUAN = int(int(_checks["lender received from the run (wei)"]["value"]) / WEI)
VB_DETAIL = _checks["vendor bond decreased by the penalty (wei)"]["detail"]  # "posted X, now Y"
VB_FROM, VB_TO = [int(int(p.split()[-1]) / WEI) for p in VB_DETAIL.split(",")]


# ---------- slide spec ----------
def T(x, y, w, h, text, size=18, color=INK, bold=False, align="left", space=6):
    return dict(k="t", x=x, y=y, w=w, h=h, text=text, size=size, color=color, bold=bold, align=align, space=space)

def R(x, y, w, h, fill, line=None):
    return dict(k="r", x=x, y=y, w=w, h=h, fill=fill, line=line)

def I(x, y, w, path, crop=None):
    im = Image.open(path)
    if crop:
        im = im.crop(crop)
    build = Path(tempfile.gettempdir()) / "zishi-deck"
    build.mkdir(exist_ok=True)
    out = build / (Path(path).stem + ("-crop" if crop else "") + ".png")
    im.save(out)
    return dict(k="i", x=x, y=y, w=w, h=w * im.height / im.width, path=str(out))

def header(kicker, title, tsize=34):
    # sits in the template banner (dark, 1.17 in tall); the FinTechathon logo takes the right 2.5 in
    return [dict(T(0.48, 0.17, 9.9, 0.3, kicker.upper(), 12, "ff6b8b", True), fixed=True),
            dict(T(0.48, 0.45, 10.0, 0.6, title, 24, "ffffff", True), fixed=True)]

def stat(x, y, w, big, label, color=RED, bsize=40, lsize=15):
    return [R(x, y, 0.06, 1.55, color), T(x + 0.25, y - 0.05, w - 0.25, 0.8, big, bsize, INK, True),
            T(x + 0.25, y + 0.75, w - 0.25, 0.9, label, lsize, GREY)]

def card(x, y, w, h, title, body, accent=RED, tsize=20, bsize=15):
    return [R(x, y, w, h, CARD, LINE), R(x, y, w, 0.07, accent),
            T(x + 0.3, y + 0.3, w - 0.6, 0.6, title, tsize, INK, True),
            T(x + 0.3, y + 0.95, w - 0.6, h - 1.1, body, bsize, GREY)]


SLIDES = []
def slide(els, notes):
    SLIDES.append(dict(els=els, notes=notes))

# 1 title
# on the template cover: project name and team name are the template's own fields, set in build_pptx
slide([T(0.74, 5.12, 6.6, 0.6, "A wallet for a drone on hire purchase. The AI participates, the money rules live in code.", 16, "ffffff", space=2),
       T(0.80, 6.1, 8.0, 0.5, f"{TRACK}  |  {DATE}\nBuilt on WeBank's FISCO BCOS and FATE.", 14, "dfe2ee", space=2),
       T(0.80, 6.65, 8.0, 0.3, f"Live demo  {URL}", 15, "ff6b8b", True)],
      "Zishi, 自食其力, means to live by your own labour. It is software that lets a machine pay for itself. "
      "Our first machine is an agricultural spraying drone on hire purchase. The AI is a real financial participant: "
      "it plans jobs, signs payment intents and repays a lender, but every money rule lives in code and on chain. (30 seconds)")

# 2 problem
slide(header("The pilot", "Skilled pilots, no collateral, drones in the showroom")
      + stat(M, 2.4, 3.9, "200,000+", "agricultural drone pilots in China")
      + stat(M + 4.0, 2.4, 3.9, "55,000 yuan", "about what one spraying drone costs")
      + stat(M + 8.0, 2.4, 3.9, "3 to 15 yuan", "earned per mu sprayed")
      + [T(M, 4.75, 11.5, 1.4, "A pilot with nothing to pledge cannot get a loan. The lender cannot see the work, so it cannot lend against it. The drone sits in a showroom. The field waits.", 22, INK)],
      "China has more than 200,000 agricultural drone pilots. A spraying drone costs about 55,000 yuan and earns 3 to 15 yuan per mu. "
      "The work is there and it pays. But a pilot with nothing to pledge cannot borrow, because the lender cannot see the work. (40 seconds)")

# 3 what zishi does
slide(header("What Zishi does", "The machine is the borrower")
      + card(M, 2.2, 3.75, 3.4, "1  Earn", "The drone earns per mu sprayed. The farmer locks the fee first and confirms the work.", RED, 22, 19)
      + card(M + 3.95, 2.2, 3.75, 3.4, "2  Repay", "Each fee splits by rule: an instalment inside a 20 to 45 percent band, a rail fee, a bond top up, the rest to the pilot.", GOLD, 22, 19)
      + card(M + 7.9, 2.2, 3.75, 3.4, "3  Own", "At zero the title moves on chain. The drone belongs to the person who flew it.", INK, 22, 19)
      + [T(M, 5.95, 11.8, 0.5, "The lease, the bond and the title belong to the drone's wallet. Its verified work is the collateral.", 18, GREY)],
      "Zishi turns the drone into the borrower. It earns per mu, confirmed by the farmer who pays. Every fee repays the lender "
      "inside a fixed band, fills a bond and leaves the rest to the pilot. At zero the title moves on chain to the pilot. (40 seconds)")

# 4 five moments
moments = "1   Job declined on economics\n2   Instalment paid to the lender\n3   Faked flight log held\n4   Slash by proof\n5   Title transferred to the pilot"
slide(header("Live demo", "One drone, five moments, about three minutes")
      + [T(M, 2.3, 4.3, 3.2, moments, 19, INK, space=14),
         T(M, 5.55, 4.2, 1.2, "Every step is a signed event on the screen. Press Start on the Fleet page.", 15, GREY),
         I(5.45, 2.3, 7.15, DOCS / "screenshots/demo.png")],
      "This is the live Fleet page. One press of Start plays five moments: a job declined because the margin is too thin, "
      "an instalment paid, a faked flight log held by the oracle, the vendor bond slashed by proof, and the title moving to the pilot. "
      "Every step is a signed event you can inspect. (50 seconds, or show the live page)")

# 5 ask the agent
def bubble(y, who, text, h, agent):
    x = M + (1.2 if agent else 0)
    return [R(x, y, 7.3, h, CARD if agent else "eef0fa", LINE), R(x, y, 0.06, h, RED if agent else GOLD),
            T(x + 0.25, y + 0.1, 6.9, 0.35, who, 12, RED if agent else "4a5bb0", True),
            T(x + 0.25, y + 0.42, 6.9, h - 0.5, text, 15, INK)]
slide(header("Ask the agent why", "The pilot can question a decision")
      + bubble(2.15, "PILOT", "Why was this job declined?", 0.95, False)
      + bubble(3.2, "AGENT", "The fee is 60 yuan (20 mu at 3 yuan per mu). After the instalment, charging and the rail fee, 22.8 yuan is left. The reserve rule needs 30. The break even is 3.7 yuan per mu.", 1.35, True)
      + bubble(4.65, "PILOT", "Lower the reserve for this job.", 0.95, False)
      + bubble(5.7, "AGENT", "I cannot. The reserve rule is the lender's floor. Changing it needs the operator key (L3, the physical button) and the lender.", 1.1, True)
      + [T(9.75, 2.2, 2.85, 4.6, "Every number and decision comes from code. The model only words it.\n\nFree models only. No key means labelled template text.\n\nEnglish or 中文.\n\nThe person affected can ask why: PIPL Article 24.", 16, GREY, space=4)],
      "After a job is declined, the pilot can ask the agent why. It answers from the engine's own numbers and offers the rate that clears, 3.7 yuan per mu. "
      "If the pilot asks it to bend a rule, it says no: that needs the operator key and the lender. The model only writes the words; "
      "every number is checked against the facts. This is also how we meet PIPL Article 24 on automated decisions. (50 seconds)")

# 6 money in code
slide(header("Money in code, not in the model", "The model drafts. Code decides the payment.")
      + [T(M, 2.25, 5.3, 2.4, "Four tiers on every signed intent\nL0 read, L1 draft: the model\nL2 pay inside caps and allowlists: the machine key\nL3 the operator's physical button", 18, INK, space=8)]
      + stat(M, 4.75, 2.6, "52 of 52", "attacks blocked", RED, 30, 13)
      + stat(M + 2.75, 4.75, 2.6, "7 of 7", "benign messages pass", GOLD, 30, 13)
      + [T(M, 6.4, 5.4, 0.6, "21 policy rule lines. A naive wallet falls to 51 of 52. The classifier alone holds 13 of 52; the rules and the policy do the rest.", 13, GREY),
         I(6.6, 2.25, 6.0, DOCS / "screenshots/evidence.png", crop=(0, 120, 2880, 1680))],
      "Money decisions are made in code. The model can read and draft; the machine key pays only inside caps and allowlists; "
      "the top tier is a physical button. 21 policy rule lines check every signed intent. 52 of 52 attacks are blocked and 7 of 7 honest messages pass. "
      "Honestly: the classifier alone stops only 13; the rules do the real work. Judges can type their own attack on the Evidence page. (55 seconds)")

# 7 bonds
slide(header("Two bonds, slashed only by proof", "The pilot never pays for the AI")
      + card(M, 2.3, 5.75, 2.6, "Pilot bond", "Earned from fees. Slashed for the pilot's own cheating, such as a faked flight log.", GOLD, 24, 20)
      + card(M + 6.05, 2.3, 5.75, 2.6, "Vendor bond, 5,000 yuan", "Posted by us. Slashed for the agent's mistakes, never the pilot's money.", RED, 24, 20)
      + [T(M, 5.3, 11.8, 1.5, "A slash needs a signed intent in the on chain log that breaks a rule the contract can check against its own records. Only a party of the job can submit. No trusted judge.", 19, INK)],
      "There are two bonds. The pilot bond pays for the pilot's cheating. The vendor bond, 5,000 yuan posted by us, pays for the agent's mistakes, "
      "so the pilot never pays for the AI. A slash needs proof: a signed intent in the log that breaks a rule the contract checks itself. No judge. (40 seconds)")

# 8 on chain
slide(header("On chain", "One loop, two ledgers, every receipt re-checked")
      + card(M, 2.3, 3.75, 3.9, f"{FISCO_CALLS}", f"contract calls on FISCO BCOS from one engine fleet run. Lender paid {LENDER_YUAN} yuan through Lease, vendor bond {VB_FROM:,} to {VB_TO:,} yuan by proof, lease at 0, title to the pilot.", RED, 34, 16)
      + card(M + 3.95, 2.3, 3.75, 3.9, "9 of 9", "contract calls in one full job on the deployed contracts on Base Sepolia, real value: fee locked, confirmed, instalment paid, vendor bond slashed, title moved.", GOLD, 34, 16)
      + card(M + 7.9, 2.3, 3.75, 3.9, "329 of 329", "receipts re-fetched by pnpm verify:all on Base Sepolia and FISCO BCOS, with 46 intent hashes decoded from calldata.", INK, 34, 16)
      + [T(M, 6.4, 11.8, 0.5, "4 contracts deployed on Base Sepolia, source verified on Sourcify. 15 of 15 intents sent to both chains in the same step, same hash.", 14, GREY)],
      f"Money does move on chain. In one engine fleet run on FISCO BCOS, {FISCO_CALLS} contract calls succeeded: the lender received {LENDER_YUAN} yuan through Lease, "
      f"the vendor bond went from {VB_FROM:,} to {VB_TO:,} by proof, and the title went to the pilot. On Base Sepolia one full job ran on the deployed contracts, 9 of 9 calls with value. "
      "And you do not have to trust us: pnpm verify:all re-fetches all 329 receipts we claim. (55 seconds)")

# 9 webank stack
slide(header("Built on WeBank's stack", "FISCO BCOS as the lender's ledger, FATE for credit")
      + card(M, 2.3, 3.75, 4.1, "FISCO BCOS 3.16.3", "The same four contracts as the lender's ledger.\n30 of 30 intent hashes read back equal the engine's.\nAgent bound IntentLogV2 fixes front running.", RED, 22, 17)
      + card(M + 3.95, 2.3, 3.75, 4.1, "FATE 2.2", "Three lenders train one pilot credit model, no raw records shared.\nAUC 0.729 alone, 0.782 federated, the same as pooling.\nNew pilot rate error 335 to 153 basis points.", GOLD, 22, 17)
      + card(M + 7.9, 2.3, 3.75, 4.1, "Model fingerprint", "The sha256 of the FATE model is committed to the FISCO IntentLog and read back.\nA lender or regulator can prove which model priced a pilot.", INK, 22, 17),
      "We built on WeBank's own stack. FISCO BCOS runs the same four contracts as the lender's ledger, and every intent hash read back matches. "
      "FATE lets three lenders learn pilot credit together without sharing rows: the lender with new pilots goes from AUC 0.729 to 0.782, the same as pooling all data. "
      "The model's fingerprint is on FISCO, so anyone can check which model set a rate. (50 seconds)")

# 10 security and compliance
sec = ("Actor keys random per server instance, in memory. The model process holds none.\n"
       "Front running fixed: IntentLogV2 accepts a commit only from the job's agent (on FISCO).\n"
       "84 Foundry tests, 15 contract layer attacks among them.\n"
       "0 open findings of medium or higher: Slither (42 triaged), pnpm audit, gitleaks.")
comp = ("PIPL Art. 24: decisions by written rules, the person can ask why. Hashes only on chain.\n"
        "NFRA: a licensed leasing company is the lessor and lender.\n"
        "PBoC: fees settle through a licensed bank or payment institution, or e-CNY. The chain records, it does not custody.\n"
        "CAAC: the pilot holds the operator licence, the drone is real name registered.")
slide(header("Security and compliance", "Built for a regulated lender")
      + card(M, 2.3, 5.75, 4.4, "Security", sec, RED, 22, 16)
      + card(M + 6.05, 2.3, 5.75, 4.4, "Compliance map", comp, GOLD, 22, 16),
      "Keys are random per server and never reach the model. The front running hole in the first intent log is fixed in IntentLogV2 on FISCO; "
      "on the Base Sepolia v1 it is still open and we say so. 84 Foundry tests, and the scans show no open finding of medium or higher. "
      "On compliance: PIPL Article 24, a licensed lessor under NFRA rules, payments through PBoC licensed rails, and CAAC licensed pilots. (50 seconds)")

# 11 measured
slide(header("Measured, not asserted", "Learned components, held out sets, 95% intervals")
      + stat(M, 2.3, 3.8, "72 of 72", "tampered flights held, interval 94.9 to 100%", RED, 34, 14)
      + stat(M + 3.95, 2.3, 3.8, "3.0%", "oracle false alarm, 13 of 428 honest, interval 1.8 to 5.1%", GOLD, 34, 14)
      + stat(M + 7.9, 2.3, 3.8, "93.1%", "classifier accuracy on 87 held out messages, recall 81.3% (64.7 to 91.1%)", INK, 34, 14)
      + stat(M, 4.55, 3.8, "25,593 to 19,739", "yuan, the lender's p95 loss without and with the two bonds", RED, 26, 14)
      + stat(M + 3.95, 4.55, 3.8, "13.8% to 0.0%", "track feature false alarms at full detection, before and after Kalman smoothing", GOLD, 26, 14)
      + stat(M + 7.9, 4.55, 3.8, "100%", "classifier precision, 26 of 26, false alarm 0 of 55", INK, 26, 14)
      + [T(M, 6.55, 11.8, 0.5, "Fixed seeds, synthetic data labelled as such. None of these components can sign, pay or move a cap.", 14, GREY)],
      "The learned parts are measured on held out sets with intervals. The oracle holds 72 of 72 tampered flights with a 3 percent false alarm rate. "
      "The classifier is right 93.1 percent of the time and never flags an honest message. The two bonds cut the lender's bad case loss from 25,593 to 19,739 yuan. "
      "None of these models can move money. (45 seconds)")

# 12 honest, next, close
real = ("Contracts deployed and verified on Base Sepolia; full jobs with value on Base and FISCO.\n"
        "Real FATE training, free models live in the public app.\n"
        "On FISCO the engine settles every payment on the contracts, with retries and refunds.\n"
        "Still limited: Base runs anchor payments to spare a small wallet; FISCO is one node of ours; "
        "spraying telemetry is synthetic, real flights only used to test the oracle.")
nxt = ("One licensed lessor pilot with ten drones.\nReal flight logs through the same oracle.\nA FISCO BCOS node run by the lessor.\n"
       "A secure element on each drone.\nThen any machine with verifiable paid work: delivery drones, robot arms, robotaxis.")
slide(header("Honest about what is real", "What runs today, and what comes next")
      + card(M, 2.1, 5.75, 3.8, "Real today", real, RED, 22, 15)
      + card(M + 6.05, 2.1, 5.75, 3.8, "Next", nxt, GOLD, 22, 15)
      + [R(0, 6.1, W, 0.8, INK),
         T(M, 6.25, 6.2, 0.5, "zishi-fintechathon.vercel.app", 19, "ffffff", True),
         T(M + 5.6, 6.3, 6.25, 0.5, "github.com/KarinaKKarinaK/shenzhen-fintechathon", 15, "ffffff", False, "right")],
      "To be clear about what is real: the contracts are deployed and verified, and full jobs moved value on both chains. On FISCO the engine itself settles every payment on the contracts and recovers when a call fails; "
      "on Base we anchor payments to spare a small test wallet, FISCO is our own node, and spraying telemetry is synthetic. Next is one licensed lessor with ten drones and real flight logs. "
      "The drone is the first machine; the same rail fits any machine whose paid work can be verified. Try it at the live demo. Thank you. (50 seconds)")


# ---------- pptx on the template ----------
def build_pptx(out):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    prs = Presentation(TEMPLATE)
    cover, toc, sample, thanks = list(prs.slides)
    ids = prs.slides._sldIdLst

    def drop(slide):
        for sid in list(ids):
            if prs.part.related_part(sid.rId) is slide.part:
                prs.part.drop_rel(sid.rId); ids.remove(sid)

    def put(sl, e, dy=0.0):
        y = e["y"] + (0 if e.get("fixed") else dy)
        if e["k"] == "r":
            sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(e["x"]), Inches(y), Inches(e["w"]), Inches(e["h"]))
            sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(e["fill"])
            if e["line"]:
                sh.line.color.rgb = RGBColor.from_string(e["line"]); sh.line.width = Pt(0.75)
            else:
                sh.line.fill.background()
            sh.shadow.inherit = False
        elif e["k"] == "i":
            sl.shapes.add_picture(e["path"], Inches(e["x"]), Inches(y), Inches(e["w"]), Inches(e["h"]))
        else:
            tb = sl.shapes.add_textbox(Inches(e["x"]), Inches(y), Inches(e["w"]), Inches(e["h"]))
            tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            for j, line in enumerate(e["text"].split("\n")):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[e["align"]]
                p.space_after = Pt(e["space"]); p.line_spacing = 1.1
                r = p.add_run(); r.text = line
                f = r.font; f.size = Pt(e["size"]); f.bold = e["bold"]; f.name = FONT
                f.color.rgb = RGBColor.from_string(e["color"])

    def retext(shape, text, size=None, width=None):
        # keep the template's run formatting, swap the words
        p = shape.text_frame.paragraphs[0]
        runs = p.runs
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        if size:
            runs[0].font.size = Pt(size)
        if width:
            shape.width = Inches(width)

    # cover: the template's two fields (作品名 project name, 团队名 team name) plus track, date, demo link
    title_box, team_box = list(cover.shapes)
    retext(title_box, "Zishi 自食其力")
    title_box.top = Inches(4.5)
    retext(team_box, f"Team {TEAM}", 18)
    team_box.top = Inches(5.68)
    first = SLIDES[0]
    for e in first["els"]:
        put(cover, e)
    cover.notes_slide.notes_text_frame.text = first["notes"]

    # contents: the template's six numbered entries
    boxes = {sh.text_frame.text: [] for sh in toc.shapes}
    for sh in toc.shapes:
        boxes[sh.text_frame.text].append(sh)
    for sh, label in zip(boxes["目录内容"], CONTENTS):
        retext(sh, label, 22, 2.6)
        sh.top = sh.top + Inches(0.08)
    toc.notes_slide.notes_text_frame.text = "Six parts: the pilot, how it works, the live demo, safety, on chain, proof and next. (5 seconds, or skip)"

    layout = next(l for l in prs.slide_layouts if l.name == "标题和内容")
    n = len(SLIDES) + 2
    for i, s in enumerate(SLIDES[1:], start=2):
        sl = prs.slides.add_slide(layout)
        for e in s["els"]:
            put(sl, e, DY)
        put(sl, T(M, H - 0.45, 4.0, 0.3, "Zishi 自食其力", 11, GREY))
        put(sl, T(W - M - 2.0, H - 0.45, 2.0, 0.3, f"{i + 1} / {n}", 11, GREY, align="right"))
        sl.notes_slide.notes_text_frame.text = s["notes"]

    # THANKS last, with the demo link
    put(thanks, T(0.85, 5.9, 8.0, 0.4, f"Live demo  {URL}", 18, "ffffff", True))
    thanks.notes_slide.notes_text_frame.text = "Thank you. Questions welcome; the live demo is open at " + URL + "."
    for sid in list(ids):
        if prs.part.related_part(sid.rId) is thanks.part:
            ids.remove(sid); ids.append(sid)
    drop(sample)  # last, so new slide part names do not clash with the template's
    prs.save(out)


# ---------- pdf / png: the built pptx drawn as HTML, printed by the system Chrome ----------
# No office renderer here (Keynote scripting fails on this machine), so each slide is redrawn from the pptx itself:
# the layout's background picture from the template, then every shape at its exact position, size, font size and colour.
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"


def pptx_html(pptx, only=None):
    import base64, html
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.enum.dml import MSO_COLOR_TYPE
    prs = Presentation(pptx)
    px = 96 / 914400  # emu to css px
    uri = lambda blob, ct: f"data:{ct};base64,{base64.b64encode(blob).decode()}"
    pages = []
    for i, sl in enumerate(prs.slides):
        if only is not None and i != only:
            continue
        divs = []
        shapes = [sh for sh in sl.slide_layout.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE] + list(sl.shapes)
        for sh in shapes:
            pos = f"left:{sh.left*px:.1f}px;top:{sh.top*px:.1f}px;width:{sh.width*px:.1f}px;height:{sh.height*px:.1f}px"
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                divs.append(f'<img src="{uri(sh.image.blob, sh.image.content_type)}" style="{pos}">')
            elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                fill = f"#{sh.fill.fore_color.rgb}" if sh.fill.type == 1 else "transparent"
                line = f"box-shadow:inset 0 0 0 1px #{sh.line.color.rgb};" if sh.line.fill.type == 1 else ""
                divs.append(f'<div style="{pos};background:{fill};{line}"></div>')
            elif sh.has_text_frame:
                tf = sh.text_frame
                nowrap = "white-space:nowrap;" if tf.word_wrap is False or tf._txBody.bodyPr.get("wrap") == "none" else ""
                ps = []
                for para in tf.paragraphs:
                    spans = []
                    for r in para.runs:
                        f = r.font
                        c = f"#{f.color.rgb}" if f.color.type == MSO_COLOR_TYPE.RGB else "#ffffff"
                        size = (f.size.pt if f.size else 18) * 4 / 3
                        spans.append(f'<span style="font-size:{size:.1f}px;color:{c};font-weight:{700 if f.bold else 400}">{html.escape(r.text)}</span>')
                    sa = para.space_after.pt * 4 / 3 if para.space_after is not None else 0
                    al = {2: "center", 3: "right"}.get(para.alignment, "left") if para.alignment else "left"
                    ps.append(f'<p style="margin:0 0 {sa:.1f}px 0;text-align:{al}">{"".join(spans) or "&nbsp;"}</p>')
                divs.append(f'<div class="t" style="{pos};{nowrap}">{"".join(ps)}</div>')
        pages.append("<section>" + "".join(divs) + "</section>")
    css = (f"@page{{size:{W*96:.0f}px {H*96:.0f}px;margin:0}}html,body{{margin:0;background:#fff}}"
           f"section{{position:relative;width:{W*96:.0f}px;height:{H*96:.0f}px;overflow:hidden;page-break-after:always}}"
           f"section>*{{position:absolute;box-sizing:border-box}}.t{{font-family:{FONT},'PingFang SC',sans-serif;line-height:1.15}}")
    return f'<!doctype html><html><head><meta charset="utf-8"><style>{css}</style></head><body>{"".join(pages)}</body></html>'


def chrome(*args):
    subprocess.run([CHROME, "--headless=new", "--user-data-dir=/tmp/zishi-chrome", "--disable-gpu", "--hide-scrollbars", *args],
                   check=True, capture_output=True)


def main():
    pptx = DOCS / "Zishi_Pitch_Deck.pptx"; pdf = DOCS / "pdf/Zishi_Pitch_Deck.pdf"
    build_pptx(pptx)
    tmp = Path(tempfile.gettempdir()) / "zishi-deck"; tmp.mkdir(exist_ok=True)
    page = tmp / "deck.html"; page.write_text(pptx_html(pptx))
    chrome("--no-pdf-header-footer", f"--print-to-pdf={pdf}", page.as_uri())
    if "--png" in sys.argv:
        out = Path(sys.argv[sys.argv.index("--png") + 1]); out.mkdir(parents=True, exist_ok=True)
        subprocess.run(["pdftoppm", "-png", "-r", "96", str(pdf), str(out / "slide")], check=True)
    print(f"{len(SLIDES) + 2} slides: docs/Zishi_Pitch_Deck.pptx, docs/pdf/Zishi_Pitch_Deck.pdf")


if __name__ == "__main__":
    main()
