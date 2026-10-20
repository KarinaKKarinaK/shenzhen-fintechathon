#!/usr/bin/env python3
"""Fills the organisers' official .docx templates (fintechathon.g-ican.com/templates, copies in
docs/template/) from the repository's Markdown, keeping the template's cover, header table, headings
and guiding bullets. Each "### <bullet text>" section in the Markdown replaces the "Write here" line
under the matching bullet.

Run from web/:  python3 scripts/officialdocs.py
Writes: docs/Zishi_Security_Self_Assessment.docx, docs/Zishi_Execution_Evidence.docx, docs/Zishi_Demo_Video_Plan.docx
"""
import copy
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[2]
TEAM = {"Team name": "Zishi", "Competition track": "International AI Track, Topic F: Agent Wallet", "Submission date": "20 October 2026"}
TEAM.update({"Team": TEAM["Team name"], "Track": TEAM["Competition track"], "Date": TEAM["Submission date"]})


def runs(par, text, size=None):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Menlo"; r.font.size = Pt(8.5)
        else:
            r = par.add_run(part)
        if size and not part.startswith("`"):
            r.font.size = Pt(size)


def blocks(lines):
    """Markdown lines to ("p", text) / ("table", rows) / ("li", text) blocks."""
    out, i = [], 0
    while i < len(lines):
        l = lines[i]
        if l.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[-| :]+\|$", lines[i].strip()):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            out.append(("table", rows))
            continue
        if re.match(r"^[-*] ", l):
            out.append(("li", l[2:]))
        elif re.match(r"^\d+\. ", l):
            out.append(("li", re.sub(r"^\d+\. ", "", l)))
        elif l.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i]); i += 1
            out.append(("code", "\n".join(code)))
        elif l.strip() and not l.startswith("#"):
            out.append(("p", l))
        i += 1
    return out


def fill(template: Path, md: Path, out: Path):
    doc = Document(template)
    for t in doc.tables:
        for row in t.rows:
            key = row.cells[0].text.strip().split(" / ")[-1]
            if key in TEAM:
                cell = row.cells[-1]
                cell.paragraphs[0].text = ""
                cell.paragraphs[0].add_run(TEAM[key])
    text = md.read_text(encoding="utf8").splitlines()
    sections, cur = {}, None
    for l in text:
        m = re.match(r"^### (.+)", l)
        if m:
            cur = m.group(1).strip().lower(); sections[cur] = []; continue
        if l.startswith("## "):
            cur = None; continue
        if cur is not None:
            sections[cur].append(l)
    paras = doc.paragraphs
    used = set()
    for idx, p in enumerate(paras):
        key = p.text.strip().lower()
        if p.style.name.startswith("List") and key in sections:
            target = next((q for q in paras[idx + 1: idx + 3] if q.text.strip() == "Write here"), None)
            if target is None:
                continue
            anchor = target._p
            for kind, val in blocks(sections[key]):
                if kind == "table":
                    ncol = len(val[0])
                    tbl = doc.add_table(rows=len(val), cols=ncol)
                    tbl.style = "Table Grid"
                    for r, row in enumerate(val):
                        for c in range(ncol):
                            cp = tbl.cell(r, c).paragraphs[0]
                            runs(cp, row[c] if c < len(row) else "", 8)
                            if r == 0:
                                for rr in cp.runs: rr.bold = True
                    anchor.addnext(tbl._tbl); anchor = tbl._tbl
                    spacer = doc.add_paragraph(); anchor.addnext(spacer._p); anchor = spacer._p
                else:
                    np_ = doc.add_paragraph(style="List Bullet" if kind == "li" else None)
                    if kind == "code":
                        r = np_.add_run(val); r.font.name = "Menlo"; r.font.size = Pt(8)
                    else:
                        runs(np_, val, 10)
                    anchor.addnext(np_._p); anchor = np_._p
            target._p.getparent().remove(target._p)
            used.add(key)
    missing = [k for k in sections if k not in used]
    if missing:
        raise SystemExit(f"{md.name}: sections with no matching template bullet: {missing}")
    doc.save(out)
    print(f"wrote {out.relative_to(ROOT)} from {md.relative_to(ROOT)} on {template.name}")


fill(ROOT / "docs/template/international-security-self-assessment-en.docx", ROOT / "SELF_ASSESSMENT.md", ROOT / "docs/Zishi_Security_Self_Assessment.docx")
fill(ROOT / "docs/template/international-execution-evidence-en.docx", ROOT / "docs/EXECUTION_EVIDENCE_SUMMARY.md", ROOT / "docs/Zishi_Execution_Evidence.docx")
fill(ROOT / "docs/template/demo-video-guide-en.docx", ROOT / "docs/DEMO_VIDEO_GUIDE.md", ROOT / "docs/Zishi_Demo_Video_Plan.docx")
